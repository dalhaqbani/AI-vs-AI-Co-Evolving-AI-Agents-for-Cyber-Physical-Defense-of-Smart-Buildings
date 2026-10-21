import argparse
import sys
from collections import Counter

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Same map as the cleaning notebook: 10 raw labels -> the 4 categories the watchdog trains on.
COLLAPSE_MAP = {
    "DDoS-TCP_Flood": "flood",
    "DDoS-SYN_Flood": "flood",
    "DDoS-UDP_Flood": "flood",
    "DoS-TCP_Flood": "flood",
    "DoS-SYN_Flood": "flood",
    "DoS-UDP_Flood": "flood",
    "DNS_Spoofing": "spoofing",
    "MITM-ArpSpoofing": "spoofing",
    "CommandInjection": "command_injection",
    "BenignTraffic": "benign",
}
CATEGORIES = ["benign", "flood", "spoofing", "command_injection"]

parser = argparse.ArgumentParser()
parser.add_argument("csv_path")
parser.add_argument("--label-col", help="Label column name. Default: first column containing 'label'.")
parser.add_argument("--chunksize", type=int, default=1_000_000)
parser.add_argument("--out", default="ciciot2023_benchmark_profile.xlsx")
args = parser.parse_args()

# Read only the header first, then only the label column, so memory stays small.
header = pd.read_csv(args.csv_path, nrows=0).columns.tolist()
label_col = args.label_col
if label_col is None:
    found = [c for c in header if "label" in c.lower()]
    if not found:
        sys.exit("No column containing 'label' found. Columns: " + ", ".join(header))
    label_col = found[0]
print("Label column:", label_col)

raw_counts = Counter()
rows_read = 0
for chunk in pd.read_csv(args.csv_path, usecols=[label_col], chunksize=args.chunksize):
    labels = chunk[label_col].astype(str).str.strip()
    raw_counts.update(labels.value_counts().to_dict())
    rows_read += len(chunk)
    print(f"  read {rows_read:,} rows", flush=True)

print("\nRaw label counts:")
for k, v in raw_counts.most_common():
    print(f"  {k}: {v:,}")

unknown = sorted(set(raw_counts) - set(COLLAPSE_MAP))
if unknown:
    sys.exit("\nLabels not in COLLAPSE_MAP (send these to Mariam, then add them to the map): "
             + ", ".join(unknown))

# ---------- workbook ----------
F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F3864")
body = Font(name=F, size=10)
inp = Font(name=F, size=10, color="0000FF")
bold = Font(name=F, size=10, bold=True)
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = hdr_font, hdr_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


wb = Workbook()
cc = wb.active
cc.title = "Class counts"
cc.append(["Raw label", "Rows", "Collapsed category", "Share of total"])
raw_order = list(COLLAPSE_MAP)
for lab in raw_order:
    cc.append([lab, raw_counts.get(lab, 0), COLLAPSE_MAP[lab]])
cc["A12"] = "Total rows"
cc["B12"] = "=SUM(B2:B11)"
for r in range(2, 12):
    cc[f"D{r}"] = f"=B{r}/$B$12"
cc["A14"] = "Collapsed category"
cc["B14"] = "Rows"
cc["C14"] = "Share of total"
for i, cat in enumerate(CATEGORIES):
    r = 15 + i
    cc[f"A{r}"] = cat
    cc[f"B{r}"] = f"=SUMIF($C$2:$C$11,A{r},$B$2:$B$11)"
    cc[f"C{r}"] = f"=B{r}/$B$12"
header(cc, 1, 4)
header(cc, 14, 3)
for r in list(range(1, 13)) + list(range(14, 19)):
    for c in range(1, 5 if r < 13 else 4):
        cell = cc.cell(row=r, column=c)
        if cell.value is None:
            continue
        cell.border = box
        if cell.font != hdr_font:
            cell.font = body
for r in range(2, 12):
    cc[f"B{r}"].font = inp
    cc[f"B{r}"].number_format = "#,##0"
    cc[f"D{r}"].number_format = "0.00%;-0.00%;-"
for c in ("A12", "B12"):
    cc[c].font = bold
cc["B12"].number_format = "#,##0"
for r in range(15, 19):
    cc[f"B{r}"].number_format = "#,##0"
    cc[f"C{r}"].number_format = "0.00%;-0.00%;-"
cc.column_dimensions["A"].width = 24
cc.column_dimensions["B"].width = 16
cc.column_dimensions["C"].width = 22
cc.column_dimensions["D"].width = 16

sm = wb.create_sheet("Summary", 0)
sm.append(["Metric (collapsed 4 categories)", "CICIoT2023"])
CR = "'Class counts'!$B$15:$B$18"
CN = "'Class counts'!$A$15:$A$18"
metrics = [
    ("Total rows", "='Class counts'!B12", "#,##0"),
    ("Benign rows", "='Class counts'!B15", "#,##0"),
    ("Attack rows", "=B2-B3", "#,##0"),
    ("Benign share", "=B3/B2", "0.00%"),
    ("Attack share", "=B4/B2", "0.00%"),
    ("Largest class (rows)", f"=MAX({CR})", "#,##0"),
    ("Largest class (name)", f"=INDEX({CN},MATCH(B7,{CR},0))", "@"),
    ("Smallest class (rows)", f"=MIN({CR})", "#,##0"),
    ("Smallest class (name)", f"=INDEX({CN},MATCH(B9,{CR},0))", "@"),
    ("Imbalance ratio (largest / smallest)", '=IF(B9=0,"n/a: a class has no rows",B7/B9)', "#,##0.0"),
    ("Flag: benign class", '=IF(B5<0.05,"Benign under 5%: resample or weight","OK")', "@"),
    ("Flag: smallest class", '=IF(B9<100,"Under 100 rows: score unreliable","OK")', "@"),
]
for name, formula, fmt in metrics:
    sm.append([name, formula])
header(sm, 1, 2)
for i, (_, _, fmt) in enumerate(metrics):
    r = 2 + i
    sm.cell(row=r, column=1).font = bold
    sm.cell(row=r, column=1).border = box
    c = sm.cell(row=r, column=2)
    c.font, c.border, c.number_format = body, box, fmt
    c.alignment = Alignment(horizontal="center" if fmt == "@" else "right", wrap_text=True, vertical="center")
sm.column_dimensions["A"].width = 38
sm.column_dimensions["B"].width = 30
sm.row_dimensions[12].height = 30
sm.row_dimensions[13].height = 30

nt = wb.create_sheet("Notes")
notes = [
    "Source and assumptions",
    f"Counts come from ciciot2023_filtered_clean.csv ({rows_read:,} rows), read in chunks of {args.chunksize:,} rows using the '{label_col}' column only.",
    "Raw label counts on the Class counts sheet are typed inputs (blue). Every other number is a formula that recalculates from them.",
    "The 10 raw labels collapse into 4 categories using the same map as the cleaning notebook: 6 flood types -> flood, DNS_Spoofing and MITM-ArpSpoofing -> spoofing, CommandInjection -> command_injection, BenignTraffic -> benign.",
    "All shares, flags and the imbalance ratio on the Summary sheet are computed on the collapsed 4, which is what the watchdog (Layer 1) trains on.",
    "Imbalance ratio = largest collapsed class / smallest collapsed class.",
    "CICIoT2023 is network flow data, so it benchmarks the network watchdog (Layer 1) only. The physical-consistency checker (Layer 2) uses the synthetic sensor-value attacks.",
]
for n in notes:
    nt.append([n])
nt["A1"].font = Font(name=F, bold=True, size=11)
for r in range(2, len(notes) + 1):
    nt.cell(row=r, column=1).font = body
    nt.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
nt.column_dimensions["A"].width = 110
wb.save(args.out)
print("\nSaved", args.out)
